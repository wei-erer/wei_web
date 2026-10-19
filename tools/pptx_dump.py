# -*- coding: utf-8 -*-
"""把 PPTX 的所有文字、表格、图片信息导出成便于阅读的文本。"""
import sys
from pptx import Presentation
from pptx.util import Emu

INFILE = sys.argv[1]
OUTFILE = sys.argv[2]

prs = Presentation(INFILE)
lines = []
lines.append(f"文件: {INFILE}")
lines.append(f"幻灯片尺寸: {prs.slide_width} x {prs.slide_height} EMU")
lines.append(f"幻灯片数量: {len(prs.slides)}")
lines.append("")

def dump_shape(shape, indent="  "):
    if shape.has_text_frame:
        txt = shape.text_frame.text.strip()
        if txt:
            lines.append(indent + "[TEXT] " + txt.replace("\n", " / "))
    if getattr(shape, "has_table", False) and shape.has_table:
        lines.append(indent + "[TABLE]")
        for row in shape.table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            lines.append(indent + "  | " + " | ".join(cells) + " |")
    if shape.shape_type == 13:  # picture
        try:
            lines.append(indent + f"[PICTURE] name={shape.name} w={Emu(shape.width).inches:.2f}in h={Emu(shape.height).inches:.2f}in")
        except Exception:
            lines.append(indent + "[PICTURE]")
    if getattr(shape, "has_chart", False) and shape.has_chart:
        lines.append(indent + "[CHART] " + shape.chart.chart_type)
    if shape.shape_type == 6:  # group
        for sub in shape.shapes:
            dump_shape(sub, indent + "    ")

for i, slide in enumerate(prs.slides, 1):
    lines.append(f"===== Slide {i} =====")
    for shape in slide.shapes:
        dump_shape(shape)
    if slide.has_notes_slide:
        note = slide.notes_slide.notes_text_frame.text.strip()
        if note:
            lines.append("  [NOTES] " + note.replace("\n", " / "))
    lines.append("")

with open(OUTFILE, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("slides:", len(prs.slides))
print("written:", OUTFILE)
